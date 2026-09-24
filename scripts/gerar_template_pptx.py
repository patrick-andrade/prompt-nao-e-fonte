#!/usr/bin/env python3
"""Gera o template PPTX de referência (entrada da Dimensão 3).

Adaptação visual da OECD Economic Outlook 2026/1, sem marca nem fotografia.
Uso (raiz 2026/):

    uv run python scripts/gerar_template_pptx.py
"""

from __future__ import annotations

import shutil
import subprocess
import sys
import tempfile
import zipfile
from copy import deepcopy
from pathlib import Path
from xml.etree import ElementTree as ET

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
DESTINO = ROOT / "03-relatorio-qmd" / "template-referencia.pptx"

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS = {"a": A}

# Office theme: dk/lt + 6 accents + hyperlinks
SCHEME = {
    "dk1": ("sys", "windowText", "0C2440"),
    "lt1": ("sys", "window", "FFFFFF"),
    "dk2": ("srgb", None, "243A50"),
    "lt2": ("srgb", None, "EAF4FB"),
    "accent1": ("srgb", None, "176FC1"),
    "accent2": ("srgb", None, "49A9DF"),
    "accent3": ("srgb", None, "119DA4"),
    "accent4": ("srgb", None, "8CA9C0"),
    "accent5": ("srgb", None, "637D94"),
    "accent6": ("srgb", None, "CBE8F9"),
    "hlink": ("srgb", None, "176FC1"),
    "folHlink": ("srgb", None, "135D9B"),
}


def _utf8_stdio() -> None:
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8")
        except (AttributeError, OSError):
            pass


def _quarto_pandoc() -> list[str]:
    quarto = shutil.which("quarto")
    if quarto:
        return [quarto, "pandoc"]
    return ["pandoc"]


def dump_reference_pptx(path: Path) -> None:
    cmd = _quarto_pandoc() + [
        "-o",
        str(path),
        "--print-default-data-file",
        "reference.pptx",
    ]
    subprocess.run(cmd, check=True, cwd=ROOT)


def _set_color(parent: ET.Element, kind: str, sys_val: str | None, hex_rgb: str) -> None:
    for child in list(parent):
        parent.remove(child)
    if kind == "sys":
        el = ET.SubElement(parent, f"{{{A}}}sysClr")
        el.set("val", sys_val or "windowText")
        el.set("lastClr", hex_rgb)
    else:
        el = ET.SubElement(parent, f"{{{A}}}srgbClr")
        el.set("val", hex_rgb)


def patch_theme(pptx_path: Path) -> None:
    with zipfile.ZipFile(pptx_path, "r") as zin:
        nomes = zin.namelist()
        temas = [n for n in nomes if n.startswith("ppt/theme/") and n.endswith(".xml")]
        if not temas:
            raise FileNotFoundError("Nenhum ppt/theme/*.xml no PPTX de referência.")
        dados = {n: zin.read(n) for n in nomes}

    for tema in temas:
        root = ET.fromstring(dados[tema])
        scheme = root.find(".//a:clrScheme", NS)
        if scheme is None:
            continue
        scheme.set("name", "AzulRelatorio2026")
        for tag, (kind, sys_val, hex_rgb) in SCHEME.items():
            node = scheme.find(f"a:{tag}", NS)
            if node is None:
                node = ET.SubElement(scheme, f"{{{A}}}{tag}")
            _set_color(node, kind, sys_val, hex_rgb)
        dados[tema] = ET.tostring(root, encoding="utf-8", xml_declaration=True)

    tmp = pptx_path.with_suffix(".patched.pptx")
    with zipfile.ZipFile(tmp, "w", compression=zipfile.ZIP_DEFLATED) as zout:
        for nome in nomes:
            zout.writestr(nome, dados[nome])
    tmp.replace(pptx_path)


def _set_placeholder_font(shape, *, size: int, color: str, bold: bool = False) -> None:
    paragraph = shape.text_frame.paragraphs[0]
    paragraph.font.name = "Aptos Display" if bold else "Aptos"
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = RGBColor.from_string(color)
    # Pandoc substitui os parágrafos de exemplo. A formatação precisa estar
    # também nos níveis do placeholder para os novos parágrafos a herdarem.
    styles = shape.text_frame._txBody.find(f"{{{A}}}lstStyle")
    for level in range(1, 10):
        properties = styles.find(f"{{{A}}}lvl{level}pPr")
        if properties is None:
            properties = OxmlElement(f"a:lvl{level}pPr")
            styles.append(properties)
        previous = properties.find(f"{{{A}}}defRPr")
        if previous is not None:
            properties.remove(previous)
        properties.append(deepcopy(paragraph._p.pPr.defRPr))


def _add_layout_bar(layout, x: float, y: float, width: float, height: float) -> None:
    """Copia um retângulo de slide para a árvore do layout, atrás dos placeholders."""
    scratch = Presentation()
    shape = scratch.slides.add_slide(scratch.slide_layouts[6]).shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(73, 169, 223)
    shape.line.fill.background()
    element = deepcopy(shape._element)
    element.xpath("./p:nvSpPr/p:cNvPr")[0].set("id", str(100 + len(layout.shapes)))
    layout.shapes._spTree.insert(2, element)


def style_layouts(pptx_path: Path) -> None:
    deck = Presentation(pptx_path)
    for index, layout in enumerate(deck.slide_layouts):
        layout.background.fill.solid()
        layout.background.fill.fore_color.rgb = RGBColor.from_string(
            "0C2440" if index == 0 else "FFFFFF"
        )
        for shape in layout.shapes:
            if not shape.is_placeholder:
                continue
            original_left, original_width = shape.left, shape.width
            kind = str(shape.placeholder_format.type)
            if "SUBTITLE" in kind:
                _set_placeholder_font(shape, size=20, color="CBE8F9")
                shape.top, shape.height = Inches(3.25), Inches(1.8)
            elif "TITLE" in kind:
                _set_placeholder_font(
                    shape, size=36 if index == 0 else 26,
                    color="FFFFFF" if index == 0 else "0C2440", bold=True,
                )
                shape.top = Inches(1.15 if index == 0 else 0.15)
                shape.height = Inches(1.8 if index == 0 else 0.9)
            elif "FOOTER" in kind or "DATE" in kind or "SLIDE_NUMBER" in kind:
                _set_placeholder_font(
                    shape, size=9, color="CBE8F9" if index == 0 else "637D94"
                )
            elif shape.has_text_frame:
                _set_placeholder_font(shape, size=19, color="243A50")
                shape.top, shape.height = Inches(1.35), Inches(4.0)
            # Criar uma transformação local só com top/height zera x/cx.
            # Preservar também a posição horizontal e a largura herdadas.
            shape.left = original_left if original_left is not None else Inches(0.5)
            shape.width = original_width if original_width is not None else Inches(9.0)
        if index == 0:
            _add_layout_bar(layout, 0, 0, 0.18, 5.625)
        else:
            _add_layout_bar(layout, 0.50, 1.12, 9.0, 0.055)
    deck.save(pptx_path)


def main() -> int:
    _utf8_stdio()
    DESTINO.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as td:
        bruto = Path(td) / "reference.pptx"
        dump_reference_pptx(bruto)
        patch_theme(bruto)
        style_layouts(bruto)
        shutil.copyfile(bruto, DESTINO)
    print(f"STATUS: template PPTX em {DESTINO.relative_to(ROOT).as_posix()}")
    print("STATUS: layouts azul-marinho/branco com acentos azuis (sem números fiscais).")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
